from io import BytesIO
from pathlib import Path
from subprocess import CalledProcessError, run
from urllib.parse import urljoin

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import Workbook

from apps.main import tests as main_tests

from .importer import read_curve
from .models import ICSetting


class ICManagementTests(TestCase):
    def test_live_search_and_sorting(self):
        run(["node", str(Path(__file__).parent / "testdata/ic_list.cjs")], check=True)
        for kind in ("rx", "charger"):
            response = self.client.get(reverse("ICmanage:ic-list", args=[kind]), {"q": "<missing>"})
            self.assertEqual(len(response.context["items"]), ICSetting.objects.filter(kind=kind).count())
            self.assertContains(response, 'value="&lt;missing&gt;"')
            self.assertContains(response, 'ICmanage/ic-list.js')
            self.assertContains(response, 'id="ic-sort"', count=1)
            self.assertContains(response, '最後修改時間：新到舊')
            self.assertContains(response, '<th scope="col">最後修改時間</th>', html=True)
            self.assertNotContains(response, 'class="sort-button"')

    def test_import_time_survives_edit_and_updated_time_is_displayed(self):
        from datetime import timedelta

        from django.utils import timezone

        item = ICSetting.objects.create(kind="rx", name="時間測試", model_number="TEST", mode="fixed", efficiency=80)
        created = item.created_at
        self.assertIsNotNone(created)
        ICSetting.objects.filter(pk=item.pk).update(updated_at=created - timedelta(days=1))
        item.name = "時間測試更新"
        item.save()
        item.refresh_from_db()
        self.assertEqual(item.created_at, created)
        self.assertGreaterEqual(item.updated_at, created)
        response = self.client.get(reverse("ICmanage:ic-list", args=["rx"]))
        self.assertContains(response, timezone.localtime(item.updated_at).strftime('%Y-%m-%d %H:%M:%S'))
        self.assertContains(response, f'data-updated="{int(item.updated_at.timestamp())}"')

    def test_import_is_the_only_creation_entry(self):
        before = ICSetting.objects.count()
        for kind in ("rx", "charger"):
            page = self.client.get(reverse("ICmanage:ic-list", args=[kind]))
            self.assertContains(page, reverse("ICmanage:ic-import", args=[kind]))
            self.assertNotContains(page, "新增命名設定")
            self.assertEqual(self.client.get(f"/ics/{kind}/new/").status_code, 404)
            self.assertEqual(self.client.post(f"/ics/{kind}/new/", self.data()).status_code, 404)
        self.assertNotContains(self.client.get(reverse("main:index")), "digDatabaseSave")
        self.assertEqual(ICSetting.objects.count(), before)

    def test_fixed_efficiency_display_rounds_without_changing_stored_value(self):
        item = ICSetting.objects.get(code="cps4019")
        item.mode, item.efficiency, item.points = "fixed", 92.35, []
        item.save()
        self.assertContains(self.client.get(reverse("ICmanage:ic-list", args=["rx"])), "92.4%")
        item.refresh_from_db()
        self.assertEqual(item.efficiency, 92.35)

    def test_ic_status_is_not_exposed_in_management_or_catalog(self):
        for kind in ("rx", "charger"):
            for route in ("ic-list", "ic-import"):
                page = self.client.get(reverse("ICmanage:" + route, args=[kind]))
                self.assertNotContains(page, 'class="badge pending"')
                self.assertNotContains(page, 'name="provisional"')
                self.assertNotContains(page, "<th>狀態</th>")
        catalog = self.client.get(reverse("main:index")).context["curves"]
        self.assertNotIn("rxPending", catalog)
        self.assertNotIn("chargerPending", catalog)

    def test_crud_redirects_preserve_proxy_prefix(self):
        for kind in ("rx", "charger"):
            item = ICSetting.objects.filter(kind=kind).first()
            edit_url = reverse("ICmanage:ic-edit", args=[kind, item.pk])
            edited = self.client.post(edit_url, self.data())
            delete_url = reverse("ICmanage:ic-delete", args=[kind, item.pk])
            deleted = self.client.post(delete_url)
            for url, response in ((edit_url, edited), (delete_url, deleted)):
                self.assertEqual(response.status_code, 302)
                self.assertTrue(response["Location"].startswith("../"))
                for prefix in ("", "/proxy/8000"):
                    self.assertEqual(
                        urljoin(prefix + url, response["Location"]),
                        prefix + reverse("ICmanage:ic-list", args=[kind]),
                    )

    def data(self, **overrides):
        return {
            "name": "測試設定", "model_number": "TEST-01", "mode": "curve",
            "axis": "power", "points": "[[0, 60], [2, 80], [4, 90]]",
            "efficiency": "", "source": "測試條件",
            **overrides,
        }

    def test_seed_and_separate_lists(self):
        self.assertEqual(ICSetting.objects.count(), 12)
        for kind, own, other in [("rx", "CPS4019", "MP2733GQC-C04L"), ("charger", "MP2733GQC-C04L", "CPS4019")]:
            response = self.client.get(reverse("ICmanage:ic-list", args=[kind]))
            self.assertContains(response, own)
            self.assertNotContains(response, other)
        self.assertEqual(self.client.get("/ics/unknown/").status_code, 404)

    def test_named_crud_and_database_lookup(self):
        for kind in ("rx", "charger"):
            existing = ICSetting.objects.create(kind=kind, name="待編輯", model_number="TEST",
                                                mode="fixed", efficiency=80)
            response = self.client.post(reverse("ICmanage:ic-edit", args=[kind, existing.pk]), self.data())
            self.assertEqual(response.status_code, 302)
            item = ICSetting.objects.get(kind=kind, name="測試設定")
            self.assertEqual(item.points, [[0, 60], [2, 80], [4, 90]])
            self.assertContains(self.client.get(reverse("main:index")), item.code)
            field = "rxIc" if kind == "rx" else "chargerIc"
            payload = {"state": {**main_tests.BASE_STATE, field + "Select": item.code, "sysPower": 3}}
            result = main_tests.calculate_js(payload["state"], self.client.get(reverse("main:index")).context["curves"])
            self.assertEqual(result[field + "Eff"], 85)
            self.assertNotIn("暫定", result["rxTip" if kind == "rx" else "chargerTip"])
            edit_url = reverse("ICmanage:ic-edit", args=[kind, item.pk])
            self.assertEqual(self.client.post(edit_url, self.data(name="已改名", mode="fixed", efficiency="77", points="")).status_code, 302)
            item.refresh_from_db()
            self.assertEqual(item.name, "已改名")
            self.assertEqual(item.points, [])
            result = main_tests.calculate_js(payload["state"], self.client.get(reverse("main:index")).context["curves"])
            self.assertEqual(result[field + "Eff"], 77)
            self.assertNotIn("暫定", result["rxTip" if kind == "rx" else "chargerTip"])
            delete_url = reverse("ICmanage:ic-delete", args=[kind, item.pk])
            self.assertEqual(self.client.get(delete_url).status_code, 200)
            self.assertTrue(ICSetting.objects.filter(pk=item.pk).exists())
            self.assertEqual(self.client.post(delete_url).status_code, 302)
            self.assertNotIn(item.code, self.client.get(reverse("main:index")).context["curves"][kind + "Fixed"])

    def test_invalid_data_and_duplicate_names(self):
        item = ICSetting.objects.get(code="cps4019")
        url = reverse("ICmanage:ic-edit", args=["rx", item.pk])
        invalid = [
            {"points": "[[0,70],[0,80],[2,90]]"},
            {"points": "[[0,70],[1,101],[2,90]]"},
            {"points": "[[0,70],[1,80]]"},
            {"points": "null"}, {"points": "not json"},
            {"points": "[[false,70],[1,80],[2,90]]"},
            {"mode": "fixed", "efficiency": "NaN"},
            {"mode": "fixed", "efficiency": "101"},
            {"mode": "fixed", "efficiency": ""},
            {"name": ""}, {"axis": "unknown"},
        ]
        for values in invalid:
            with self.subTest(values=values):
                response = self.client.post(url, self.data(**values))
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.context["form"].errors)
        self.assertEqual(ICSetting.objects.count(), 12)
        self.client.post(url, self.data())
        other = ICSetting.objects.filter(kind="rx").exclude(pk=item.pk).first()
        response = self.client.post(reverse("ICmanage:ic-edit", args=["rx", other.pk]), self.data())
        self.assertIn("name", response.context["form"].errors)
        self.assertEqual(ICSetting.objects.filter(name="測試設定").count(), 1)

    def test_anonymous_access_csrf_and_kind_isolation(self):
        item = ICSetting.objects.get(code="cps4019")
        urls = [reverse("ICmanage:ic-edit", args=["rx", item.pk]), reverse("ICmanage:ic-delete", args=["rx", item.pk])]
        for url in urls:
            self.assertEqual(self.client.get(url).status_code, 200)
        self.assertEqual(self.client.post(reverse("ICmanage:ic-delete", args=["charger", item.pk])).status_code, 404)
        secure = Client(enforce_csrf_checks=True)
        for url in urls:
            self.assertEqual(secure.post(url, self.data()).status_code, 403)
        self.assertEqual(ICSetting.objects.count(), 12)

    def test_database_overrides_original_catalog_and_zero_fixed(self):
        item = ICSetting.objects.get(code="cps4019")
        self.client.post(reverse("ICmanage:ic-edit", args=["rx", item.pk]), self.data(mode="fixed", efficiency="0", points=""))
        payload = {"state": {**main_tests.BASE_STATE, "rxIcSelect": "cps4019"}}
        response = main_tests.calculate_js(payload["state"], self.client.get(reverse("main:index")).context["curves"])
        self.assertEqual(response["rxIcEff"], 0)
        self.assertIsNone(response["pInTotal"])

    def test_ma_curve_matches_ampere_curve_after_save_and_reload(self):
        for kind in ("rx", "charger"):
            existing = ICSetting.objects.filter(kind=kind).first()
            url = reverse("ICmanage:ic-edit", args=[kind, existing.pk])
            response = self.client.post(url, self.data(axis="ma", points="[[100, 80], [500, 92], [1000, 95]]"))
            self.assertEqual(response.status_code, 302)
            item = ICSetting.objects.get(kind=kind, name="測試設定")
            self.assertEqual(item.axis, "ma")
            self.assertEqual(item.points[1], [500, 92])
            catalog = self.client.get(reverse("main:index")).context["curves"]
            self.assertEqual(catalog[kind][item.code]["axis"], "current")
            field = "rxIc" if kind == "rx" else "chargerIc"
            for current_ma, expected in [(100, 80), (500, 92), (750, 93.5), (1000, 95)]:
                state = {**main_tests.BASE_STATE, "batCapacity": current_ma, "batMaxC": 1,
                         field + "Select": item.code}
                result = main_tests.calculate_js(state, catalog)
                self.assertAlmostEqual(result[field + "Eff"], expected)
            for current_ma in (99, 1001, 1500):
                state = {**main_tests.BASE_STATE, "batCapacity": current_ma, "batMaxC": 1,
                         field + "Select": item.code}
                with self.assertRaises(CalledProcessError) as error:
                    main_tests.calculate_js(state, catalog)
                self.assertIn("充電電流超出允許範圍", error.exception.stderr)
            edit = self.client.get(reverse("ICmanage:ic-edit", args=[kind, item.pk]))
            self.assertEqual(edit.context["form"].initial["axis"], "ma")
            item.refresh_from_db()
            self.assertEqual(item.points[1], [500, 92])


class ICImportTests(TestCase):
    def test_excel_import_uses_displayed_efficiency(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Charger_CPS5201"
        sheet.append(["Iout (mA)", "效率（小數）"])
        for current, value, fmt in [(100, 0.943, "0.0000"), (101, 0.94309, "0.0000"),
                                    (102, 0.9431799999999999, "0.0000"), (103, 0.94325, "0.0000"),
                                    (104, 0.9435, "0.0%")]:
            sheet.append([current, value])
            sheet.cell(sheet.max_row, 2).number_format = fmt
        data = BytesIO()
        workbook.save(data)
        result = read_curve(SimpleUploadedFile("curve.xlsx", data.getvalue()), sheet.title, 5)
        self.assertEqual(result["points"], [[100, 94.3], [101, 94.31], [102, 94.32], [103, 94.33], [104, 94.4]])
        response = self.client.post(reverse("ICmanage:ic-import", args=["charger"]), self.payload(
            "", model_number=sheet.title, file=SimpleUploadedFile("curve.xlsx", data.getvalue())))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(ICSetting.objects.get(name="CPS5201_5V").points, result["points"])

    def test_names_are_generated_from_model_and_voltage(self):
        url = reverse("ICmanage:ic-import", args=["charger"])
        self.assertNotContains(self.client.get(url), 'name="name"')
        csv = "current_ma,efficiency_percent\n100,80\n500,90\n1000,95"
        for voltage in (5, 12):
            response = self.client.post(url, self.payload(
                csv, model_number=f"Charger_MP2733_{voltage}V", voltage=str(voltage), name="ignored"))
            self.assertEqual(response.status_code, 302)
            item = ICSetting.objects.get(name=f"MP2733_{voltage}V")
            self.assertEqual(item.model_number, "MP2733")
            self.assertIn(f"Charger_MP2733_{voltage}V", item.source)
        before = ICSetting.objects.count()
        response = self.client.post(url, self.payload(csv, model_number="MP2733"))
        self.assertIn("__all__", response.context["form"].errors)
        response = self.client.post(url, self.payload(csv, model_number="X" * 120))
        self.assertIn("model_number", response.context["form"].errors)
        self.assertEqual(ICSetting.objects.count(), before)

    def upload(self, text):
        return SimpleUploadedFile("curve.csv", text.encode("utf-8-sig"))

    def payload(self, text, **kwargs):
        return {"model_number": "TEST", "voltage": "5",
                "source": "實測", "file": self.upload(text), **kwargs}

    def test_csv_saved_and_available_without_overwriting(self):
        text = "voltage_v,current_ma,efficiency_fraction\n5,1000,0.9\n12,100,0.5\n5,100,0\n5,300,\n5,500,0.8\n"
        for kind in ("rx", "charger"):
            url = reverse("ICmanage:ic-import", args=[kind])
            self.assertContains(self.client.get(reverse("ICmanage:ic-list", args=[kind])), url)
            self.assertContains(self.client.get(url), 'multipart/form-data')
            response = self.client.post(url, self.payload(text))
            self.assertEqual(response.status_code, 302)
            self.assertEqual(urljoin("/proxy/8000" + url, response["Location"]),
                             "/proxy/8000" + reverse("ICmanage:ic-list", args=[kind]))
            item = ICSetting.objects.get(kind=kind, name="TEST_5V")
            self.assertEqual(item.points, [[100, 0], [500, 80], [1000, 90]])
            self.assertEqual(item.axis, "ma")
            self.assertIn("5V", item.source)
            page = self.client.get(reverse("main:index"))
            self.assertEqual(page.context["curves"][kind][item.code]["data"], [[0.1, 0], [0.5, 80], [1, 90]])
            response = self.client.post(url, self.payload(text))
            self.assertIn("__all__", response.context["form"].errors)
            self.assertEqual(ICSetting.objects.filter(kind=kind, name=item.name).count(), 1)

    def test_bad_uploads_never_write_and_csrf_is_required(self):
        url = reverse("ICmanage:ic-import", args=["rx"])
        header = "current_a,efficiency_percent\n"
        invalid = ["0,80\n0,90\n1,95", "0,80\n1,NaN\n2,90", "0,80\n1,101\n2,90",
                   "0,80\n1,-1\n2,90", "0,80\n1,=80+1\n2,90", "0,80\n1,90",
                   "0,\n1,\n2,", "-1,80\n1,90\n2,95", "0,80\n1,Infinity\n2,90"]
        before = ICSetting.objects.count()
        for rows in invalid:
            with self.subTest(rows=rows):
                response = self.client.post(url, self.payload(header + rows))
                self.assertEqual(response.status_code, 200)
                self.assertIn("file", response.context["form"].errors)
        for upload in [SimpleUploadedFile("bad.xlsx", b"not zip"),
                       SimpleUploadedFile("bad.txt", b"x"),
                       SimpleUploadedFile("large.csv", b"x" * (5 * 1024 * 1024 + 1))]:
            response = self.client.post(url, self.payload("", file=upload))
            self.assertIn("file", response.context["form"].errors)
        self.assertEqual(ICSetting.objects.count(), before)
        self.assertEqual(Client(enforce_csrf_checks=True).post(url, self.payload(header + "0,80\n1,85\n2,90")).status_code, 403)
        self.assertEqual(self.client.get("/ics/unknown/import/").status_code, 404)

    def test_excel_voltage_units_linear_model_and_dense_curve(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Charger_Test"
        sheet.append(["Charger test"])
        sheet.append(["理論值，非實測"])
        sheet.append(["VIN (V)", "Iout (mA)", "VBAT (V)", "效率（小數）", "Pout (W)"])
        for i in range(100, 1601):
            sheet.append([5, i, 3.8, 0.7 + min(i, 800) / 10000, i * 0.0038])
        sheet.append([12, 100, 3.8, None])
        data = BytesIO()
        workbook.save(data)
        def upload():
            return SimpleUploadedFile("curve.xlsx", data.getvalue())
        result = read_curve(upload(), "Charger_Test", 5)
        self.assertEqual(result["count"], 1501)
        self.assertEqual(len(result["points"]), 3)
        self.assertEqual(result["axis"], "ma")
        self.assertAlmostEqual(result["points"][1][0], 800)
        self.assertIn("理論值", result["notes"])
        for sheet_name, voltage in [("missing", 5), ("Charger_Test", 12)]:
            with self.assertRaises(ValueError):
                read_curve(upload(), sheet_name, voltage)
        url = reverse("ICmanage:ic-import", args=["charger"])
        response = self.client.post(url, self.payload("", file=upload(), model_number="Charger_Test"))
        self.assertEqual(response.status_code, 302)
        item = ICSetting.objects.get(name="Test_5V")
        self.assertIn("理論值", item.source)
        self.assertIn("Charger_Test", item.source)
        self.assertEqual(item.model_number, "Test")

    def test_combined_model_field_strips_only_category_prefix(self):
        url = reverse("ICmanage:ic-import", args=["rx"])
        page = self.client.get(url)
        self.assertContains(page, "IC 型號／Excel 工作表名稱")
        self.assertNotContains(page, 'name="worksheet"')
        csv = "current_ma,efficiency_percent\n100,80\n500,90\n1000,95"
        for entered, expected in [
            ("Rx_CV8045D", "CV8045D"), ("CHERGER_MP2733", "MP2733"),
            ("charger-Test", "Test"), ("rX Test", "Test"),
            ("MP2733", "MP2733"), ("RX123", "RX123"), ("TEST_RX_01", "TEST_RX_01"),
        ]:
            with self.subTest(entered=entered):
                response = self.client.post(url, self.payload(csv, model_number=entered))
                self.assertEqual(response.status_code, 302)
                item = ICSetting.objects.get(name=f"{expected}_5V")
                self.assertEqual(item.model_number, expected)
                item.delete()
        before = ICSetting.objects.count()
        for entered in ("Rx", "cherger_", "Charger - ", ""):
            response = self.client.post(url, self.payload(csv, model_number=entered))
            self.assertIn("model_number", response.context["form"].errors)
        self.assertEqual(ICSetting.objects.count(), before)

    def test_power_csv_and_row_limit(self):
        result = read_curve(self.upload("power_w,efficiency_percent\n0,0.5\n1,0.7\n2,0.9"), "", 5)
        self.assertEqual(result["axis"], "power")
        self.assertEqual(result["points"][0], [0, 0.5])
        with self.assertRaises(ValueError):
            read_curve(self.upload("current_a,efficiency_percent\n" + "1,80\n" * 10000), "", 5)

    def test_dense_non_collinear_curve_import_save_and_calculation(self):
        # 9999 data rows plus the header reach the existing upload row limit.
        csv = "current_ma,efficiency_percent\n" + "\n".join(
            f"{i},{80 + i % 2}" for i in range(9999)
        )
        response = self.client.post(
            reverse("ICmanage:ic-import", args=["charger"]),
            self.payload(csv, model_number="Dense"),
        )
        self.assertEqual(response.status_code, 302)
        item = ICSetting.objects.get(name="Dense_5V")
        self.assertEqual(len(item.points), 9999)
        catalog = self.client.get(reverse("main:index")).context["curves"]
        state = {**main_tests.BASE_STATE, "chargerIcSelect": item.code,
                 "batCapacity": 500.5, "batMaxC": 1}
        self.assertAlmostEqual(main_tests.calculate_js(state, catalog)["chargerIcEff"], 80.5)
        item.points.append([9999, 81])
        item.full_clean()  # Exactly 10000 points are also valid for manual editing.
        item.points.append([10000, 80])
        with self.assertRaises(ValidationError):
            item.full_clean()

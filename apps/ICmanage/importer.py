"""Read a single, explicitly selected efficiency curve without saving uploads."""

import csv
import io
import math
import re
from contextlib import contextmanager
from decimal import ROUND_HALF_UP, Decimal
from itertools import pairwise
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import number

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000
# Labels carry units: never infer percent versus fraction from the magnitude.
X_COLUMNS = {
    "Iout (mA)": ("ma", 1),
    "Iout (A)": ("current", 1),
    "current_ma": ("ma", 1),
    "current_a": ("current", 1),
    "power_w": ("power", 1),
    "Pout (W)": ("power", 1),
}
EFF_COLUMNS = {"效率（小數）": 100, "efficiency_fraction": 100,
               "效率 (%)": 1, "效率(%)": 1, "efficiency_percent": 1}
VOLT_COLUMNS = {"工作電壓 (V)", "VIN (V)", "voltage_v"}


def numeric(value, label, maximum=1_000_000):
    if isinstance(value, bool) or value is None or str(value).strip() == "":
        raise ValueError(f"{label} 不可空白，且必須是數字。")
    try:
        return number(float(value), label, 0, maximum)
    except (TypeError, ValueError) as error:
        raise ValueError(f"{label} 必須是 0–{maximum} 的有限數字。") from error


def displayed_efficiency(value, number_format, scale, label):
    """Round the efficiency as displayed, then convert to percentage points."""
    percent = number_format.endswith("%")
    numeric(value, label, 1 if percent else 100 / scale)
    decimal = Decimal(str(value))
    if number_format.lower() == "general":
        return float(decimal * scale)
    match = re.fullmatch(r"(?:0|#,##0)(?:\.([0#]{1,15}))?(%)?", number_format)
    if not match:
        raise ValueError(f"{label} 的顯示格式不支援，請改用數值或百分比格式（例如 0.0000、0.0%）。")
    places = len(match[1] or "")
    if percent:
        decimal *= 100
    displayed = decimal.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    return float(displayed * (1 if percent else scale))


def read_upload(upload):
    if upload.size > MAX_BYTES:
        raise ValueError("檔案不可超過 5 MB。")
    content = upload.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise ValueError("檔案不可超過 5 MB。")
    return content


def read_curve(upload, worksheet, voltage):
    content = read_upload(upload)
    extension = Path(upload.name).suffix.lower()
    if extension == ".csv":
        try:
            rows = csv.reader(io.StringIO(content.decode("utf-8-sig")))
            return parse_rows(rows, voltage)
        except (UnicodeDecodeError, csv.Error) as error:
            raise ValueError("CSV 必須為 UTF-8 編碼、逗號分隔。") from error
    if extension != ".xlsx":
        raise ValueError("僅支援 .xlsx 或 UTF-8 .csv。")
    with open_workbook(content) as workbook:
        if worksheet not in workbook.sheetnames:
            raise ValueError("請選擇檔案內的工作表：" + "、".join(workbook.sheetnames)[:1000])
        sheet = workbook[worksheet]
        if sheet.max_column and sheet.max_column > 100:
            raise ValueError("工作表不可超過 100 欄。")
        return parse_rows(sheet.iter_rows(), voltage, excel=True)


def read_worksheets(upload):
    if Path(upload.name).suffix.lower() != ".xlsx":
        raise ValueError("請選擇 .xlsx 檔案。")
    with open_workbook(read_upload(upload)) as workbook:
        return workbook.sheetnames


@contextmanager
def open_workbook(content):
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("Excel 解壓縮後過大，請只保留要匯入的工作表。")
            # Reject DTD/entity expansion, including installations without defusedxml.
            for entry in archive.infolist():
                if entry.filename.endswith((".xml", ".rels")):
                    xml = archive.read(entry).replace(b"\x00", b"").upper()
                    if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                        raise ValueError("不支援含 DTD 或實體宣告的 Excel。")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            yield workbook
        finally:
            workbook.close()
    except (BadZipFile, InvalidFileException, KeyError, OSError) as error:
        raise ValueError("無法讀取 Excel，請確認是有效且未加密的 .xlsx。") from error
    except Exception as error:
        # XML parser exception classes vary with the installed XML backend.
        if isinstance(error, ValueError):
            raise
        raise ValueError("Excel 格式損壞或包含不支援的內容。") from error


def parse_rows(rows, voltage, excel=False):
    points, notes = [], []
    columns = None
    skipped = 0
    for row_number, row in enumerate(rows, 1):
        if row_number > MAX_ROWS:
            raise ValueError("工作表最多可有 10000 列。")
        if len(row) > 100:
            raise ValueError(f"第 {row_number} 列超過 100 欄。")
        cells = row if excel else None
        if excel:
            row = [cell.value for cell in cells]
        if not any(value is not None and str(value).strip() for value in row):
            continue
        if columns is None:
            labels = [str(value).strip() if value is not None else "" for value in row]
            xs = [i for i, label in enumerate(labels) if label in X_COLUMNS]
            es = [i for i, label in enumerate(labels) if label in EFF_COLUMNS]
            vs = [i for i, label in enumerate(labels) if label in VOLT_COLUMNS]
            if xs and es:
                # Linear-model sheets contain both Iout and Pout; prefer Iout.
                currents = [i for i in xs if X_COLUMNS[labels[i]][0] in ("current", "ma")]
                xs = currents or xs
                if len(xs) != 1 or len(es) != 1 or len(vs) > 1:
                    raise ValueError("欄位重複，無法判定電流／功率、效率或電壓欄。")
                xi, ei = xs[0], es[0]
                vi = vs[0] if vs else None
                axis, x_scale = X_COLUMNS[labels[xi]]
                efficiency_scale = EFF_COLUMNS[labels[ei]]
                columns = (xi, ei, vi)
                continue
            notes.extend(label for label in labels if label)
            if row_number >= 10:
                break
            continue
        xi, ei, vi = columns
        if len(row) <= max(i for i in columns if i is not None):
            raise ValueError(f"第 {row_number} 列欄位不足。")
        if vi is not None and numeric(row[vi], f"第 {row_number} 列電壓") != voltage:
            continue
        if row[ei] is None or str(row[ei]).strip() == "":
            skipped += 1
            continue
        x = numeric(row[xi], f"第 {row_number} 列 X") * x_scale
        label = f"第 {row_number} 列效率"
        if excel:
            eta = displayed_efficiency(row[ei], cells[ei].number_format, efficiency_scale, label)
        else:
            numeric(row[ei], label, 100 / efficiency_scale)
            eta = float(Decimal(str(row[ei])) * efficiency_scale)
        points.append([x, eta])
    if columns is None:
        raise ValueError("前 10 列找不到含單位的電流／功率及效率欄，請參考下方格式。")
    if not points:
        if skipped:
            raise ValueError(f"{voltage}V 的效率欄全部空白，無法建立效率曲線。請提供有效效率資料；不會將空白當成 0 或自動估算。")
        raise ValueError(f"找不到 {voltage}V 的資料列，請確認匯入電壓與工作表內容一致。")
    if len(points) < 3:
        raise ValueError("此電壓至少需要 3 筆有效效率資料；空白不會轉成 0。")
    points.sort(key=lambda point: point[0])
    if any(a[0] == b[0] for a, b in pairwise(points)):
        raise ValueError("同一電壓有重複的電流／功率，請分開不同量測條件後匯入。")
    original_count = len(points)
    if original_count > 1000:
        simplified = []
        for point in points:
            while len(simplified) >= 2:
                a, b = simplified[-2:]
                expected = a[1] + (point[1] - a[1]) * (b[0] - a[0]) / (point[0] - a[0])
                if not math.isclose(b[1], expected, abs_tol=1e-10, rel_tol=0):
                    break
                simplified.pop()
            simplified.append(point)
        if len(simplified) == 2:
            simplified.insert(1, points[len(points) // 2])
        points = simplified
        # MAX_ROWS already keeps the curve within the 10000-point storage limit.
    return {"axis": axis, "points": points, "count": original_count,
            "skipped": skipped, "notes": "；".join(notes)}
